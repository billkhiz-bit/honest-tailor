"""Render a CV written in Markdown to HTML and PDF, and optionally Word.

The Markdown follows the layout described in `skills/tailor-cv/references/file-formats.md`. The HTML uses the
stylesheet beside this script (`cv-style.css`). The PDF is printed by a headless Chromium-based browser
(Edge, Chrome or Chromium), found automatically or named with --browser or the CV_BROWSER environment variable.
The Word version is for agencies and portals that ask for .docx (needs python-docx).

With no browser it still writes the HTML (and Word, when python-docx is installed). Exit code 0 means everything
asked for was written; 3 means something was not (no browser, the browser failed, or python-docx is missing for
--docx), and the output says what. Only known browser executables are ever run.

Usage: python render_cv.py "<cv>.md" [--letter] [--docx] [--out-dir <folder>] [--browser <path>]
"""
from __future__ import annotations

import argparse
import html
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
from pathlib import Path

HERE = Path(__file__).resolve().parent
STYLE = HERE / "cv-style.css"

BROWSER_CANDIDATES = [
    r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
    r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
    r"C:\Program Files\Google\Chrome\Application\chrome.exe",
    r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
    "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
    "/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge",
    "/Applications/Chromium.app/Contents/MacOS/Chromium",
]
BROWSER_COMMANDS = ["google-chrome", "google-chrome-stable", "chromium", "chromium-browser", "microsoft-edge",
                    "microsoft-edge-stable"]
# Only these executables are ever run, whatever --browser or CV_BROWSER says
BROWSER_NAMES = {"msedge.exe", "chrome.exe", "chromium.exe", "google chrome", "microsoft edge", "chromium",
                 *BROWSER_COMMANDS}
BROWSER_FLAGS = ["--headless=new", "--disable-gpu", "--no-pdf-header-footer", "--no-first-run",
                 "--disable-extensions", "--disable-background-networking"]

LINK = re.compile(r"\[([^\]]{1,300})\]\(([^)\s]{1,2000})\)")
BOLD = re.compile(r"(?<![\w*])\*\*(?=\S)(.+?)(?<=\S)\*\*(?![\w*])")
# An asterisk is italics only where it touches a word on the inside alone, so grades like "A*" stay literal
ITALIC = re.compile(r"(?<![\w*])\*(?![\s*])(.+?)(?<![\s*])\*(?![\w*])")
INLINE = re.compile(r"\[([^\]]{1,300})\]\(([^)\s]{1,2000})\)"
                    r"|(?<![\w*])\*\*(?=\S)(.+?)(?<=\S)\*\*(?![\w*])"
                    r"|(?<![\w*])\*(?![\s*])(.+?)(?<![\s*])\*(?![\w*])")
LABEL_LINE = re.compile(r"^\*\*([^*]+?:)\*\*\s*(.*)$")  # **Languages:** Python, ...
DATED_HEADER = re.compile(r"^\*\*(.+?)\*\*\s+[-–—]\s+\*(.+?)\*\s*$")  # **Employer** - *Mar 2021 – Present*
TITLED_HEADER = re.compile(r"^\*\*(.+?)\*\*\s*(.*)$")  # **Title** | stack - [link](url)
EMAIL = re.compile(r"[\w.+-]+@[\w-]+(\.[\w-]+)+")
WEB = re.compile(r"(?:[\w-]+\.)+[a-z]{2,}(?:/\S*)?", re.I)
SAFE_URL = re.compile(r"^(?:https?://|mailto:)", re.I)  # anything else (javascript:, file:) renders as plain text


def whole_italic(line):
    return len(line) > 2 and line.startswith("*") and line.endswith("*") and not line.startswith("**")


# ---------------------------------------------------------------- parsing


def parse(md_text):
    """Turn the CV Markdown into a list of blocks that both renderers understand."""
    blocks, section, entry = [], None, None

    def close_entry():
        nonlocal entry
        if entry:
            blocks.append(entry)
            entry = None

    for raw in md_text.splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# "):
            blocks.append({"type": "name", "text": line[2:].strip()})
            continue
        if line.startswith("## "):
            close_entry()
            section = line[3:].strip()
            blocks.append({"type": "section", "text": section})
            continue
        if section is None:
            # Everything between the name and the first section: contact lines, then the summary
            if "•" in line or EMAIL.search(line) or LINK.search(line):
                blocks.append({"type": "contact", "text": line})
            else:
                # A partly formatted summary keeps its Markdown, which inline_html renders (and "A*" survives)
                blocks.append({"type": "tagline", "text": line[1:-1] if whole_italic(line) else line})
            continue
        if line.startswith("- "):
            item = line[2:].strip()
            if entry:
                entry["bullets"].append(item)
            elif blocks and blocks[-1]["type"] == "bullets":
                blocks[-1]["items"].append(item)
            else:
                blocks.append({"type": "bullets", "items": [item]})
            continue
        m = LABEL_LINE.match(line)
        if m:
            close_entry()
            kind = "skill" if "SKILL" in section.upper() else "label"
            blocks.append({"type": kind, "label": m[1], "text": m[2]})
            continue
        m = DATED_HEADER.match(line)
        if m:
            close_entry()
            entry = {"type": "job", "title": m[1], "date": m[2], "role": None, "bullets": []}
            continue
        m = TITLED_HEADER.match(line)
        if m:
            close_entry()
            stack, links = m[2].strip(), ""
            if " - [" in stack:
                stack, links = stack.split(" - [", 1)
                links = "[" + links
            elif stack.startswith("- ["):
                stack, links = "", stack[2:]
            entry = {"type": "project", "title": m[1], "stack": stack.strip(), "links": links.strip(),
                     "bullets": []}
            continue
        if entry and entry["type"] == "job" and entry["role"] is None and not entry["bullets"] \
                and line.startswith("*") and not line.startswith("**"):
            # "*Analyst*" or "*Analyst* – London": either way it is the role under this job
            entry["role"] = line[1:-1] if whole_italic(line) else line
            continue
        close_entry()
        blocks.append({"type": "para", "text": line})
    close_entry()
    return blocks


def contact_parts(text):
    """Split a contact line on its bullets and say which parts are links."""
    parts = []
    for part in [p.strip() for p in text.split("•")]:
        m = LINK.fullmatch(part)
        if m:
            parts.append((m[1], m[2] if SAFE_URL.match(m[2]) else None))
        elif EMAIL.fullmatch(part):
            parts.append((part, "mailto:" + part))
        elif WEB.fullmatch(part) and not part.replace(" ", "").isdigit():
            parts.append((part, "https://" + part))
        else:
            parts.append((part, None))
    return parts


# ---------------------------------------------------------------- HTML


def _anchor(m):
    if not SAFE_URL.match(m[2]):
        return m[1]
    return f'<a href="{m[2].replace(chr(34), "&quot;")}">{m[1]}</a>'


def inline_html(text):
    out = html.escape(text, quote=False)
    # Links become placeholders first, so emphasis can never reach into an href (a "*" in a URL)
    anchors = []

    def hold(m):
        anchors.append(_anchor(m))
        return f"\x00{len(anchors) - 1}\x00"

    out = LINK.sub(hold, out)
    out = BOLD.sub(r"<strong>\1</strong>", out)
    out = ITALIC.sub(r"<em>\1</em>", out)
    return re.sub(r"\x00(\d+)\x00", lambda m: anchors[int(m[1])], out)


def html_head(title, letter):
    css = STYLE.read_text(encoding="utf-8")
    if letter:
        css = css.replace("size: A4", "size: Letter")
    lang = "en-US" if letter else "en-GB"
    return (f'<!DOCTYPE html>\n<html lang="{lang}">\n<head>\n<meta charset="UTF-8">\n'
            f"<title>{html.escape(title)}</title>\n<style>\n{css}</style>\n</head>\n")


def to_html(blocks, letter, title):
    body = []
    for b in blocks:
        t = b["type"]
        if t == "name":
            body.append(f"<h1>{html.escape(b['text'])}</h1>")
        elif t == "contact":
            parts = [f'<a href="{html.escape(url)}">{html.escape(label)}</a>' if url else html.escape(label)
                     for label, url in contact_parts(b["text"])]
            body.append('<div class="contact">\n  ' + " •\n  ".join(parts) + "\n</div>")
        elif t == "tagline":
            body.append('<div class="contact" style="font-style: italic; margin-top: -6pt;">\n  '
                        + inline_html(b["text"]) + "\n</div>")
        elif t == "section":
            body.append(f"\n<h2>{html.escape(b['text'])}</h2>")
        elif t == "skill":
            body.append(f'<div class="skills-line"><strong>{html.escape(b["label"])}</strong>'
                        f'<span class="content">{inline_html(b["text"])}</span></div>')
        elif t == "label":
            body.append(f'<div class="interests"><strong>{html.escape(b["label"])}</strong> '
                        f'{inline_html(b["text"])}</div>')
        elif t == "bullets":
            body.append(html_list(b["items"]))
        elif t == "para":
            body.append(f"<p>{inline_html(b['text'])}</p>")
        elif t == "job":
            role = f'\n  <div class="role">{inline_html(b["role"])}</div>' if b["role"] else ""
            body.append('<div class="job">\n  <div class="job-header">\n'
                        f'    <span class="job-title">{inline_html(b["title"])}</span>\n'
                        f'    <span class="date">{html.escape(b["date"])}</span>\n  </div>'
                        f"{role}\n{html_list(b['bullets'])}</div>")
        elif t == "project":
            stack = f' <span class="stack">{inline_html(b["stack"])}</span>' if b["stack"] else ""
            links = f'\n    <span class="link">{inline_html(b["links"])}</span>' if b["links"] else ""
            body.append('<div class="project">\n  <div class="project-meta">\n'
                        f'    <span><span class="project-title">{inline_html(b["title"])}</span>{stack}</span>'
                        f"{links}\n  </div>\n{html_list(b['bullets'])}</div>")
    return html_head(title, letter) + "<body>\n\n" + "\n".join(body) + "\n\n</body>\n</html>\n"


def html_list(items):
    if not items:
        return ""
    return "  <ul>\n" + "".join(f"    <li>{inline_html(i)}</li>\n" for i in items) + "  </ul>\n"


# ---------------------------------------------------------------- PDF


def _usable_browser(path):
    """The path itself if it is a known browser executable; a macOS .app bundle maps to the binary inside it."""
    path = Path(path)
    if path.suffix == ".app" and path.is_dir():
        path = path / "Contents" / "MacOS" / path.stem
    if not path.is_file() or path.name.lower() not in BROWSER_NAMES:
        return None
    return path


def find_browser(explicit=None):
    """Return a Chromium-based browser that can print to PDF, or None."""
    for choice in (explicit, os.environ.get("CV_BROWSER")):
        if choice:
            if os.name != "nt" and not any(sep in choice for sep in "/\\"):
                choice = shutil.which(choice) or choice  # a bare command name such as "chromium"
            found = _usable_browser(choice)
            if found is None:
                print(f"note: {choice!r} is not a known Edge, Chrome or Chromium executable, so it was not used")
            return found
    for path in BROWSER_CANDIDATES:
        if _usable_browser(path):
            return Path(path)
    if os.name != "nt":  # on Windows, which() also searches the current folder, so it is never used there
        for command in BROWSER_COMMANDS:
            found = shutil.which(command)
            if found and _usable_browser(found):
                return Path(found)
    return None


def print_pdf(html_path, pdf_path, browser):
    """Print the HTML to PDF. Returns (produced, page count or None)."""
    try:
        pdf_path.unlink(missing_ok=True)
    except PermissionError:
        sys.exit(f"{pdf_path.name} is open in another program. Close it and run this again.")
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as profile:
        # A throwaway profile keeps this away from any browser window that is already open
        try:
            subprocess.run([str(browser), *BROWSER_FLAGS, f"--user-data-dir={profile}",
                            f"--print-to-pdf={pdf_path}", html_path.as_uri()], capture_output=True, timeout=120)
        except (subprocess.TimeoutExpired, OSError) as e:
            print(f"PDF:  not produced: the browser failed ({type(e).__name__}).")
            return False, None
        for _ in range(40):
            if pdf_path.exists() and pdf_path.stat().st_size > 0:
                break
            time.sleep(0.25)
    if not pdf_path.exists() or pdf_path.stat().st_size == 0:
        print("PDF:  not produced: the browser ran but wrote no PDF.")
        return False, None
    try:
        from pypdf import PdfReader
        return True, len(PdfReader(str(pdf_path)).pages)
    except Exception:
        return True, None


# ---------------------------------------------------------------- Word


def to_docx(blocks, docx_path, letter):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = (Mm(215.9), Mm(279.4)) if letter else (Mm(210), Mm(297))
    sec.top_margin = sec.bottom_margin = Mm(15)
    sec.left_margin = sec.right_margin = Mm(18)
    text_width = sec.page_width - sec.left_margin - sec.right_margin
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(1)

    def link(par, text, url, size=None):
        r_id = par.part.relate_to(
            url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
            is_external=True)
        hyper = OxmlElement("w:hyperlink")
        hyper.set(qn("r:id"), r_id)
        run, rpr = OxmlElement("w:r"), OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "0563C1")
        rpr.append(color)
        if size:
            sz = OxmlElement("w:sz")
            sz.set(qn("w:val"), str(int(size * 2)))
            rpr.append(sz)
        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "single")
        rpr.append(underline)
        run.append(rpr)
        t = OxmlElement("w:t")
        t.text = text
        t.set(qn("xml:space"), "preserve")
        run.append(t)
        hyper.append(run)
        par._p.append(hyper)

    def runs(par, text, size=None, bold=False, italic=False):
        pos = 0
        for m in INLINE.finditer(text):
            if m.start() > pos:
                plain(par, text[pos:m.start()], size, bold, italic)
            if m[1] and SAFE_URL.match(m[2]):
                link(par, m[1], m[2], size)
            elif m[1]:
                plain(par, m[1], size, bold, italic)
            elif m[3]:
                plain(par, m[3], size, True, italic)
            else:
                plain(par, m[4], size, bold, True)
            pos = m.end()
        if pos < len(text):
            plain(par, text[pos:], size, bold, italic)

    def plain(par, text, size=None, bold=False, italic=False):
        r = par.add_run(text)
        r.bold, r.italic = bold or None, italic or None
        if size:
            r.font.size = Pt(size)
        return r

    def bottom_border(par):
        bdr, bottom = OxmlElement("w:pBdr"), OxmlElement("w:bottom")
        for k, v in (("val", "single"), ("sz", "8"), ("space", "1"), ("color", "000000")):
            bottom.set(qn(f"w:{k}"), v)
        bdr.append(bottom)
        par._p.get_or_add_pPr().insert_element_before(
            bdr, "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
            "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
            "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
            "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
            "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")

    def bullet_list(items):
        for item in items:
            par = doc.add_paragraph(style="List Bullet")
            par.paragraph_format.space_after = Pt(2)
            runs(par, item)

    def header(left_runs, right_text=None, right_italic=False, right_is_markdown=False, wrap=False):
        par = doc.add_paragraph()
        par.paragraph_format.space_before = Pt(6)
        par.paragraph_format.keep_with_next = True
        par.paragraph_format.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT)
        left_runs(par)
        if right_text:
            if wrap:
                # Too long for one line: Word would run the links into the stack, so they get their own line
                par.add_run().add_break()
            plain(par, "\t")
            if right_is_markdown:
                runs(par, right_text, size=10)
            else:
                plain(par, right_text, size=10, italic=right_italic)
        return par

    for b in blocks:
        t = b["type"]
        if t == "name":
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.space_after = Pt(4)
            plain(par, b["text"], size=22, bold=True)
        elif t in ("contact", "tagline"):
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.space_after = Pt(4 if t == "contact" else 8)
            if t == "tagline":
                runs(par, b["text"], size=9.5, italic=True)
                continue
            for i, (label, url) in enumerate(contact_parts(b["text"])):
                if i:
                    plain(par, " • ", size=9.5)
                link(par, label, url, 9.5) if url else plain(par, label, size=9.5)
        elif t == "section":
            par = doc.add_paragraph()
            bottom_border(par)
            par.paragraph_format.space_before = Pt(10)
            par.paragraph_format.space_after = Pt(4)
            par.paragraph_format.keep_with_next = True
            plain(par, b["text"].upper(), size=11, bold=True)
        elif t == "skill":
            par = doc.add_paragraph()
            pf = par.paragraph_format
            pf.left_indent, pf.first_line_indent = Pt(110), Pt(-110)
            pf.space_before = pf.space_after = Pt(2)
            pf.tab_stops.add_tab_stop(Pt(110))
            plain(par, b["label"], bold=True)
            plain(par, "\t")
            runs(par, b["text"])
        elif t == "label":
            par = doc.add_paragraph()
            par.paragraph_format.space_before = Pt(8)
            plain(par, b["label"] + " ", bold=True)
            runs(par, b["text"])
        elif t == "bullets":
            bullet_list(b["items"])
        elif t == "para":
            runs(doc.add_paragraph(), b["text"])
        elif t == "job":
            header(lambda p: runs(p, b["title"], bold=True), b["date"], right_italic=True)
            if b["role"]:
                par = doc.add_paragraph()
                par.paragraph_format.keep_with_next = True
                runs(par, b["role"], size=10, italic=True)
            bullet_list(b["bullets"])
        elif t == "project":
            def left(p, b=b):
                runs(p, b["title"], bold=True)
                if b["stack"]:
                    plain(p, " ")
                    runs(p, b["stack"], size=10)
            one_line = LINK.sub(r"\1", f'{b["title"]} {b["stack"]} {b["links"]}')
            header(left, b["links"] or None, right_is_markdown=True, wrap=len(one_line) > 105)
            bullet_list(b["bullets"])
    try:
        doc.save(str(docx_path))
    except PermissionError:
        sys.exit(f"{docx_path.name} is open in another program. Close it and run this again.")


# ---------------------------------------------------------------- main


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser(description="Render a CV in Markdown to HTML, PDF and optionally Word.")
    ap.add_argument("markdown", type=Path, help="the CV .md file")
    ap.add_argument("--letter", action="store_true", help="US Letter paper (US and Canada); A4 otherwise")
    ap.add_argument("--docx", action="store_true", help="also write a Word version (needs python-docx)")
    ap.add_argument("--out-dir", type=Path, help="where to write the files (default: beside the .md)")
    ap.add_argument("--browser", help="path to Edge, Chrome or Chromium (default: found automatically)")
    args = ap.parse_args()

    md = args.markdown.resolve()
    if not md.exists():
        sys.exit(f"Not found: {md}")
    out_dir = (args.out_dir or md.parent).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)

    blocks = parse(md.read_text(encoding="utf-8-sig"))
    html_path = out_dir / (md.stem + ".html")
    html_path.write_text(to_html(blocks, args.letter, md.stem), encoding="utf-8")
    print(f"HTML: {html_path}")
    complete = True
    browser = find_browser(args.browser)
    want_docx = args.docx
    if browser is None:
        # No PDF here, so still leave something the user can submit: the HTML prints to PDF from any browser
        print("PDF:  not produced: no Edge, Chrome or Chromium found (pass --browser <path to the browser's "
              "executable> or set CV_BROWSER). Open the HTML in a browser and print it to PDF; the page count is "
              "unknown.")
        complete = False
        want_docx = want_docx or _docx_available()
    else:
        pdf_path = out_dir / (md.stem + ".pdf")
        produced, pages = print_pdf(html_path, pdf_path, browser)
        if produced:
            print(f"PDF:  {pdf_path} ({pages if pages is not None else '?'} pages, "
                  f"{'Letter' if args.letter else 'A4'})")
        complete = complete and produced
    if want_docx and not _docx_available():
        print("Word: not produced: python-docx is not installed (pip install python-docx).")
        complete = False
    elif want_docx:
        docx_path = out_dir / (md.stem + ".docx")
        to_docx(blocks, docx_path, args.letter)
        print(f"Word: {docx_path}")
    return 0 if complete else 3


def _docx_available():
    try:
        import docx  # noqa: F401
        return True
    except ImportError:
        return False


if __name__ == "__main__":
    sys.exit(main())
